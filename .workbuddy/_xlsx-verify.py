# -*- coding: utf-8 -*-
"""读回 xlsx 校验：结构 / 公式 / 缓存值 / 数据验证"""

import json

from openpyxl import load_workbook

P = r"D:\LZB\LIN\川西消费明细.xlsx"
OUTF = r"D:\LZB\LIN\.workbuddy\_xlsx-verify.json"

import os

wf = load_workbook(P)
wv = load_workbook(P, data_only=True)

rep = {"size": os.path.getsize(P), "sheets": wf.sheetnames}

d_f, d_v = wf["消费明细"], wv["消费明细"]
rep["detail_rows"] = []
for r in range(3, 14):
    rep["detail_rows"].append({
        "no": d_v.cell(r, 1).value,
        "note": d_v.cell(r, 2).value,
        "cat": d_v.cell(r, 3).value,
        "amt": d_v.cell(r, 4).value,
        "remark": d_v.cell(r, 5).value,
    })
rep["total_formula"] = d_f.cell(14, 4).value
rep["total_cached"] = d_v.cell(14, 4).value
rep["total_label_cached"] = d_v.cell(14, 1).value
rep["tail_cached"] = d_v.cell(14, 5).value
rep["detail_dv"] = [str(dv.sqref) + " :: " + str(dv.formula1)
                    for dv in d_f.data_validations.dataValidation]
rep["detail_filter"] = str(d_f.auto_filter.ref)

s_f, s_v = wf["分类汇总"], wv["分类汇总"]
rep["cats"] = []
for r in range(3, 10):
    rep["cats"].append({
        "name": s_v.cell(r, 1).value,
        "formula": s_f.cell(r, 2).value,
        "cached": s_v.cell(r, 2).value,
        "cnt_formula": s_f.cell(r, 3).value,
        "cnt": s_v.cell(r, 3).value,
        "pct_formula": s_f.cell(r, 4).value,
        "pct": s_v.cell(r, 4).value,
    })
rep["cat_total_formula"] = s_f.cell(10, 2).value
rep["cat_total_cached"] = s_v.cell(10, 2).value
rep["cnt_total_cached"] = s_v.cell(10, 3).value
rep["pct_total_cached"] = s_v.cell(10, 4).value
rep["share"] = s_v.cell(12, 2).value
rep["avg_formula"] = s_f.cell(13, 2).value
rep["avg_cached"] = s_v.cell(13, 2).value

# 交叉核对
rep["check_sum_of_rows"] = sum(x["amt"] or 0 for x in rep["detail_rows"])
rep["check_sum_of_cats"] = sum(x["cached"] or 0 for x in rep["cats"])
rep["check_pct_sum"] = round(sum(x["pct"] or 0 for x in rep["cats"]), 6)
rep["check_row_count"] = len(rep["detail_rows"])

with open(OUTF, "w", encoding="utf-8") as f:
    json.dump(rep, f, ensure_ascii=False, indent=2)

print("VERIFY_DONE")
